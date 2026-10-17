"""Generate the report-format presentation (.pptx).

One slide group per report section, in the report's order:
  Start -> Abstract -> Introduction -> Literature Review -> Methodology ->
  Analysis -> Result -> Discussion -> Conclusion -> References -> End

Reuses the visual theme (colours, header bar, footer, title/thanks slides)
from presentation/make_deck.py.

Run:  python presentation/make_report_deck.py
"""
from __future__ import annotations

import os
import re
import sys

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
PROJ = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(PROJ, "report"))

from make_deck import (  # noqa: E402  (theme primitives, shared with other decks)
    NAVY, TEAL, GRAY, WHITE, LIGHT,
    _content_slide, _title_slide, _thanks_slide, _image, _notes, _new_pres,
)
from make_report import load_data  # noqa: E402  (same numbers as the report)


# --------------------------------------------------------------------------
# report-deck primitives
# --------------------------------------------------------------------------
def _bullets2(slide, items, left, top, width, height, size=15, sub=13):
    tb = slide.shapes.add_textbox(Pt(left), Pt(top), Pt(width), Pt(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (text, level) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = level
        p.space_after = Pt(5)
        for run in p.runs:
            run.font.size = Pt(size if level == 0 else sub)
            run.font.name = "Calibri"
            run.font.color.rgb = GRAY if level == 0 else RGBColor(0x50, 0x50, 0x50)
            if level == 0 and text.endswith(":") and len(text) < 40:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x1E, 0x6E, 0x8C)
    return tb


def _slide_table(slide, headers, rows, left, top, width, height, font=12):
    data = [headers] + [[str(v) for v in r] for r in rows]
    gf = slide.shapes.add_table(len(data), len(headers), Pt(left), Pt(top),
                                Pt(width), Pt(height))
    tbl = gf.table
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            for run in cell.text_frame.paragraphs[0].runs:
                run.font.size = Pt(font)
                run.font.name = "Calibri"
                run.font.color.rgb = WHITE if r == 0 else GRAY
                run.font.bold = (r == 0)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (LIGHT if r % 2 == 0 else WHITE)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return tbl


def _refs_slide(pres, idx, refs):
    s = _content_slide(pres, "References", idx)
    tb = s.shapes.add_textbox(Pt(40), Pt(92), Pt(880), Pt(415))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (num, text) in enumerate(refs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(3)
        r = p.add_run()
        r.text = f"[{num}] "
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.name = "Calibri"
        r.font.color.rgb = TEAL
        r2 = p.add_run()
        r2.text = text
        r2.font.size = Pt(9)
        r2.font.name = "Calibri"
        r2.font.color.rgb = GRAY
    return s


def load_refs():
    """Read the numbered references back out of the generated report."""
    path = os.path.join(PROJ, "report", "DrugResistance_Report.docx")
    doc = Document(path)
    refs = []
    inref = False
    for p in doc.paragraphs:
        if p.style.name == "Heading 1" and p.text.strip() == "8. References":
            inref = True
            continue
        if inref:
            if p.style.name == "Heading 1":
                break
            t = p.text.strip()
            if t:
                refs.append(re.sub(r"^\[\d+\]\s*", "", t))
    return refs


# --------------------------------------------------------------------------
# content
# --------------------------------------------------------------------------
def build_spec(d):
    ds1, ds2, sh = d["ds1"], d["ds2"], d["shared"]
    ts = d["ts_mean_auc"]
    panel = ", ".join(r["gene"] for r in d["panel"][:8])

    S = []

    S.append(("bullets", "Abstract", [
        ("•  Ovarian cancer is the most lethal gynaecologic malignancy; cisplatin is first-line, but acquired resistance drives relapse.", 0),
        ("•  Integrated two independent RNA-seq datasets: GSE303889 (SKOV-3) and GSE283381 (OVCAR5), 3 vs 3 each (n = 12).", 0),
        (f"•  DESeq2 (FDR < 0.05, |log2FC| > 1): {ds1['sig']:,} DEGs (SKOV-3); {ds2['sig']:,} DEGs (OVCAR5); {sh['total']} shared.", 0),
        ("•  Enrichment converged on ECM organisation, cell-substrate adhesion and PI3K-Akt / focal-adhesion / integrin / cadherin signalling; EMT dominant.", 0),
        (f"•  ML: within-cell-line LOOCV ROC-AUC 0.889-1.000; cross-dataset transfer partial and asymmetric (strict leak-free mean AUC {ts:.2f}).", 0),
        ("•  30-gene candidate panel (MAL2, GALNT6, ...) nominated for follow-up; exploratory given n = 12.", 0),
    ], 60, 120, 840, 380, None))

    S.append(("bullets", "Introduction", [
        ("Introduction", 0),
        ("•  Ovarian cancer: most patients present advanced; first-line = cytoreductive surgery + platinum (cisplatin) +/- taxane.", 0),
        ("•  Platinum forms DNA crosslinks -> apoptosis; intact DNA-damage response is central to cytotoxicity.", 0),
        ("•  Most patients relapse with platinum-resistant disease; resistance is the principal obstacle to survival.", 0),
        ("", 0),
        ("Problem of acquired resistance", 0),
        ("•  Adaptive, polyclonal process: reduced uptake / increased efflux, enhanced DNA repair, altered apoptosis, epigenetic and metabolic rewiring, EMT / stemness.", 0),
        ("•  Mechanisms differ between models and patients -> single-dataset signatures do not reproduce.", 0),
        ("•  Need consistently dysregulated genes across independent models as robust biomarkers and targets.", 0),
    ], 60, 115, 840, 385, None))

    S.append(("bullets", "Introduction (cont.)", [
        ("Objectives", 0),
        ("•  Identify DEGs between cisplatin-resistant and sensitive ovarian cancer cells with DESeq2.", 0),
        ("•  Validate DEGs across two independent GEO datasets -> a consistent core signature.", 0),
        ("•  Characterise affected biology via GO and KEGG enrichment.", 0),
        ("•  Build supervised classifiers and evaluate leakage-free cross-validation.", 0),
        ("•  Rank a compact, interpretable candidate biomarker panel; deliver a reproducible pipeline + dashboard.", 0),
        ("", 0),
        ("Scope and contributions", 0),
        ("•  Computational, hypothesis-generating study on public data; exploratory, not clinical validation.", 0),
        ("•  Cross-dataset core signature; nested (leakage-free) feature selection; honest external generalisation reporting.", 0),
    ], 60, 115, 840, 385, None))

    S.append(("bullets", "Literature Review", [
        ("Mechanisms", 0),
        ("•  Cisplatin: intra/inter-strand DNA crosslinks -> DNA-damage response -> apoptosis; resistance at every step (CTR1/SLC31A1 uptake, ABC efflux, DNA repair, apoptosis).", 0),
        ("•  EMT and ECM remodelling: CDH1 (E-cadherin) and EpCAM loss; collagen and MMP up-regulation; CDH1 promoter hypermethylation; reduced miR-29.", 0),
        ("•  Epigenetics: DNA-methylation drivers of acquired resistance; KDM5 histone demethylases maintain a drug-tolerant state.", 0),
        ("•  Metabolic rewiring: targetable de novo pyrimidine dependency in platinum-resistant disease (underpins the OVCAR5 model).", 0),
    ], 60, 115, 840, 385, None))

    S.append(("bullets", "Literature Review (cont.)", [
        ("Emerging themes and gap", 0),
        ("•  Inflammation / immune evasion: IL1A-IL1B inflammasome activation; tryptophan-kynurenine axis (KYNU -> AhR) suppresses anti-tumour immunity.", 0),
        ("•  Validated markers: NDRG1 (hub gene; knockdown inhibits resistant SKOV-3 proliferation); SERPINE1 (invasion/metastasis).", 0),
        ("•  Convergent pathways: PI3K-Akt, MAPK, focal adhesion / integrin, cadherin, p53.", 0),
        ("•  Cross-dataset validation and directional consistency increase confidence; single-dataset signatures are unstable.", 0),
        ("•  ML pitfalls: selection bias and preprocessing leakage in n << p settings.", 0),
        ("•  Gap: no single reproducible pipeline combining cross-dataset DEG validation + interpretable enrichment + leakage-controlled ML with honest external reporting.", 0),
    ], 60, 115, 840, 385, None))

    S.append(("table", "Methodology - Data Sources", None, None, None, None, None))

    S.append(("bullets", "Methodology - Pipeline", [
        ("Preprocessing and QC", 0),
        ("•  Gene-level counts + metadata per dataset; low-count filter rowSums >= 10; DESeq2 size-factor normalisation; log2(CPM + 1).", 0),
        ("•  Duplicate symbols collapsed (rowsum); datasets intersected -> 57,357 shared genes x 12 samples; PCA / clustering QC; seed 42.", 0),
        ("Differential expression", 0),
        ("•  DESeq2, design ~ condition, negative-binomial GLM, Wald test; sensitive = reference.", 0),
        ("•  Significance: BH FDR < 0.05 and |log2 fold change| > 1.", 0),
        ("Cross-dataset and enrichment", 0),
        ("•  Intersect DEG sets -> up / down / mixed; Spearman correlation of shared log2FC.", 0),
        ("•  clusterProfiler enrichGO (BP) and enrichKEGG (hsa), BH-adjusted p < 0.05.", 0),
    ], 60, 112, 840, 390, None))

    S.append(("bullets", "Methodology - Machine Learning", [
        ("Feature matrix and models", 0),
        ("•  Full 57,357-gene matrix (log2 CPM + 1), 12 samples; no DEG pre-filtering (avoids test leakage).", 0),
        ("•  Ridge logistic regression (glmnet), random forest (ranger, min.node.size = 1), radial SVM (e1071).", 0),
        ("Validation and leakage control", 0),
        ("•  Within-cell-line LOOCV (5 train / 1 test); cross-dataset external (train one line, predict the other).", 0),
        ("•  Feature selection nested inside each training fold (top-K by |t-statistic|, K in {20,50,100,200,500}).", 0),
        ("•  Test normalisation: target_z (batch-corrected) vs train_stats (strictly non-transductive, leak-free).", 0),
        ("Metrics and reproducibility", 0),
        ("•  Accuracy, precision, recall, F1, ROC-AUC; biomarkers ranked by cross-model selection frequency.", 0),
        ("•  R pipeline + FastAPI dashboard; deterministic (two runs byte-identical).", 0),
    ], 60, 110, 840, 392, None))

    S.append(("bullets", "Analysis - Dataset 1 (SKOV-3)", [
        (f"•  GSE303889: SKOV-3 parental vs CDDP-resistant (KDM5-associated).", 0),
        (f"•  Genes tested: {d['ds1_full']:,}", 0),
        (f"•  Significant DEGs: {ds1['sig']:,}  ({ds1['up']:,} up / {ds1['down']:,} down).", 0),
        ("•  Broad dispersion of up- and down-regulated genes -> extensive reprogramming.", 0),
        ("•  Heatmap separates sensitive vs resistant replicates cleanly.", 0),
        ("•  CDH1 (E-cadherin) down among top changes -> EMT-like shift.", 0),
    ], 40, 118, 430, 375, None,
     "imgs", [("results/dataset1/volcano_plot.png", 500, 110, 430, 390)]))

    S.append(("bullets", "Analysis - Dataset 2 (OVCAR5)", [
        ("•  GSE283381: OVCAR5 Cis-S vs Cis-R (de novo pyrimidine dependency).", 0),
        (f"•  Genes tested: {d['ds2_full']:,}", 0),
        (f"•  Significant DEGs: {ds2['sig']:,}  ({ds2['up']:,} up / {ds2['down']:,} down).", 0),
        ("•  Top up-regulated genes include IL1A and IL1B (large log2FC) -> inflammatory component.", 0),
        ("•  Magnitude and composition differ from Dataset 1 (distinct resistance mechanism).", 0),
    ], 40, 118, 430, 375, None,
     "imgs", [("results/dataset2/volcano_plot.png", 500, 110, 430, 390)]))

    S.append(("bullets", "Analysis - Cross-Dataset Overlap", [
        (f"•  {sh['total']} genes differentially expressed in BOTH datasets (core signature).", 0),
        (f"-  {sh['up']} up in both; {sh['down']} down in both; {sh['mixed']} opposite direction.", 1),
        ("•  Consistently regulated genes define a background-independent signature.", 0),
        ("•  Spearman correlation of shared log2FC ~ -0.04.", 0),
        ("•  Low correlation is biologically meaningful: the two models acquire resistance through different programmes.", 0),
    ], 40, 118, 430, 375, None,
     "imgs", [("results/comparison/venn_diagram.png", 490, 105, 450, 200),
              ("results/comparison/lfc_correlation.png", 490, 305, 450, 190)]))

    S.append(("bullets", "Analysis - Functional Enrichment", [
        ("GO (shared DEGs)", 0),
        ("•  Cell-substrate adhesion; extracellular matrix organisation; extracellular structure organisation; ossification; axon guidance.", 0),
        ("KEGG (shared DEGs)", 0),
        ("•  PI3K-Akt signalling; focal adhesion / integrin; cadherin signalling; cytoskeleton / ECM-receptor interaction; Hippo signalling.", 0),
        ("•  ECM, adhesion, survival and cytoskeletal remodelling converge on resistance -- hallmark of EMT.", 0),
    ], 40, 118, 400, 375, None,
     "imgs", [("results/enrichment/GO_dotplot_Shared_DEGs.png", 430, 100, 260, 395),
              ("results/enrichment/KEGG_dotplot_Shared_DEGs.png", 700, 100, 260, 395)]))

    S.append(("bullets", "Analysis - Machine Learning", [
        ("•  Within-cell-line LOOCV: Accuracy 0.833-1.000, ROC-AUC 0.889-1.000.", 0),
        ("-  Near-complete separation, but likely optimistic (cell-line background is highly separable).", 1),
        ("•  External, target_z (batch-corrected):", 0),
        ("-  SKOV3 -> OVCAR5: AUC 1.0 at K = 20-50.", 1),
        ("-  OVCAR5 -> SKOV3: AUC ~ 0.44-0.78 (near chance).", 1),
        (f"•  External, train_stats (strict, leak-free): mean AUC {ts:.2f}.", 0),
        ("•  Transfer is partial, strongly asymmetric, and collapses as K grows (n << p overfitting).", 0),
    ], 40, 118, 430, 375, None,
     "imgs", [("results/ml/ROC_LOOCV.png", 490, 105, 450, 200),
              ("results/ml/ROC_external.png", 490, 305, 450, 190)]))

    S.append(("bullets", "Analysis - Feature Importance", [
        ("•  Ranked by cross-dataset selection frequency across models x K.", 0),
        (f"•  Top candidates: {panel}, ...", 0),
        ("•  MAL2 and GALNT6 most consistently selected, then LDOC1, SH2D4A, FBXL21P.", 0),
        ("•  Selection frequency dominated by the SKOV3 -> OVCAR5 direction.", 0),
        ("•  Best regarded as candidates for follow-up, not a validated directional signature.", 0),
    ], 40, 118, 430, 375, None,
     "imgs", [("results/ml/feature_importance.png", 490, 100, 450, 400)]))

    panel_rows = [[r["rank"], r["gene"], r["n_select_slots"],
                   (f"{float(r['mean_abs_lfc']):.1f}"
                    if r["mean_abs_lfc"] not in ("", "NA", None) else "NA")]
                  for r in d["panel"][:10]]
    S.append(("table", "Result - Differential Expression Summary", None, None, None, None, None,
               "ds_summary"))
    S.append(("table", "Result - Candidate Biomarker Panel", None, None, None, None, None,
               "panel", panel_rows))

    S.append(("bullets", "Discussion", [
        ("Biological interpretation", 0),
        ("•  Shared signature = loss of epithelial adhesion (CDH1, EPCAM) + gain of ECM/mesenchymal components (COL1A1/5A2/6A1, MMP14) -> canonical EMT.", 0),
        ("•  OVCAR5 inflammatory axis (IL1A/IL1B); NDRG1 and KYNU provide literature-anchored nodes; SERPINE1 links to ECM remodelling and metastasis.", 0),
        ("Proposed mechanism", 0),
        ("•  Cisplatin -> DNA damage + chronic inflammation + immune-evasive metabolism -> apoptosis evasion + ECM remodelling -> integrin/focal adhesion -> Wnt/beta-catenin -> EMT -> stable resistance (PI3K-Akt as survival node).", 0),
    ], 60, 115, 840, 385, None))

    S.append(("bullets", "Discussion (cont.)", [
        ("ML interpretation", 0),
        ("•  LOOCV near-perfect but reflects intrinsic cell-line separability (sanity check, not deployable).", 0),
        (f"•  Cross-dataset transfer partial, asymmetric; strict leak-free mean AUC {ts:.2f}.", 0),
        ("Literature, limitations and future work", 0),
        ("•  Findings align with established EMT / ECM / PI3K-Akt / integrin resistance biology; low cross-dataset correlation matches known heterogeneity.", 0),
        ("•  Limitations: n = 12 (exploratory); n << p; normalisation-dependent external results; cell lines miss tumour heterogeneity; no wet-lab validation.", 0),
        ("•  Future: add datasets/replicates; benchmark batch correction; pathway-level features; validate MAL2, GALNT6, NDRG1 and the EMT axis.", 0),
    ], 60, 115, 840, 385, None))

    S.append(("bullets", "Conclusion", [
        ("•  Reproducible end-to-end transcriptomic + ML analysis of cisplatin resistance across two independent RNA-seq datasets.", 0),
        ("•  Cross-dataset validation isolated a core signature of consistently regulated genes.", 0),
        ("•  Enrichment converged on ECM remodelling, adhesion and PI3K-Akt / focal-adhesion / integrin / cadherin pathways, with EMT dominant.", 0),
        ("•  Classifiers separate phenotypes within a cell line, but cross-dataset transfer is partial, asymmetric and weakens under leak-free normalisation.", 0),
        ("•  A 30-gene candidate panel (MAL2, GALNT6, ...) plus NDRG1 nominated for follow-up.", 0),
        ("•  Deliverables: reproducible pipeline, publication-quality figures and an interactive dashboard, adaptable to other cancers/drugs/datasets.", 0),
    ], 60, 115, 840, 385, None))

    return S


# --------------------------------------------------------------------------
# render
# --------------------------------------------------------------------------
def _ds_table(slide, d):
    ds1, ds2 = d["ds1"], d["ds2"]
    _slide_table(
        slide,
        ["Dataset", "Cell line", "Genes tested", "Significant DEGs", "Up", "Down"],
        [["GSE303889", "SKOV-3", f"{d['ds1_full']:,}", f"{ds1['sig']:,}",
          f"{ds1['up']:,}", f"{ds1['down']:,}"],
         ["GSE283381", "OVCAR5", f"{d['ds2_full']:,}", f"{ds2['sig']:,}",
          f"{ds2['up']:,}", f"{ds2['down']:,}"]],
        left=120, top=170, width=720, height=150, font=15)


def _data_sources_table(slide, d):
    _slide_table(
        slide,
        ["Attribute", "GSE303889 (SKOV-3)", "GSE283381 (OVCAR5)"],
        [["Resistance model", "CDDP (chronic cisplatin; KDM5)", "Cis-R (repeated; pyrimidine dep.)"],
         ["Groups", "3 sensitive vs 3 resistant", "3 sensitive vs 3 resistant"],
         ["Platform", "Illumina NovaSeq 6000", "Illumina NextSeq 500"],
         ["Samples", "6", "6"]],
        left=70, top=150, width=820, height=180, font=14)


def _panel_table(slide, panel_rows):
    _slide_table(
        slide,
        ["Rank", "Gene", "Selected (of 15 slots)", "Mean |log2FC|"],
        panel_rows,
        left=210, top=120, width=540, height=360, font=14)


def main():
    d = load_data()
    refs = load_refs()
    spec = build_spec(d)

    pres = _new_pres()
    idx = 1
    _title_slide(
        pres,
        "Transcriptomic Profiling of Drug-Resistant Cancer Cells",
        "Cisplatin resistance in ovarian cancer: differential expression, pathways & machine-learning biomarkers",
        "Team of 3  |  Data & QC  .  DEGs & Biology  .  ML & Biomarkers",
    )
    idx = 2

    for entry in spec:
        kind, title = entry[0], entry[1]
        if kind == "table":
            tag = entry[7] if len(entry) > 7 else None
            s = _content_slide(pres, title, idx)
            if tag == "panel":
                _panel_table(s, entry[8])
            elif tag == "ds_summary":
                _ds_table(s, d)
            else:
                _data_sources_table(s, d)
        elif kind == "bullets":
            _, title, items, left, top, width, height = entry[:7]
            s = _content_slide(pres, title, idx)
            _bullets2(s, items, left, top, width, height)
            if len(entry) > 8 and entry[8] == "imgs":
                for rel, l, t, mw, mh in entry[9]:
                    _image(s, rel, l, t, mw, mh)
        idx += 1

    chunk = 10
    for start in range(0, len(refs), chunk):
        part = refs[start:start + chunk]
        numbered = [(start + i + 1, txt) for i, txt in enumerate(part)]
        _refs_slide(pres, idx, numbered)
        idx += 1

    _thanks_slide(pres)

    out = os.path.join(HERE, "DrugResistance_ReportDeck.pptx")
    pres.save(out)
    print(f"wrote {out} ({len(pres.slides)} slides)")


if __name__ == "__main__":
    main()
